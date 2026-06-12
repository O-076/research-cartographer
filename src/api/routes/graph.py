"""Graph query endpoints and WebSocket delta stream."""

import logging

from io import BytesIO
from typing import Any

from fastapi import APIRouter, HTTPException, Request, WebSocket, WebSocketDisconnect, Query, Response

from src.api.models import (
    GraphResponse, PaperStatusResponse, PIPELINE_PROGRESS,
    EdgeDetailResponse, ConsensusExplainResponse,
    VerificationResultItem, VerificationResponse,
    TracePathNode, TracePathEdge, TraceResponse,
    ReviewSection, LiteratureReviewData, LiteratureReviewResponse, ReviewDownloadRequest,
)

logger = logging.getLogger(__name__)

def build_apa_docx(review: dict[str, Any]) -> bytes:
    """Convert a review dict to an APA-formatted .docx file.

    Requires python-docx (python-docx==1.1.2 in requirements.txt).
    APA 7th edition: Times New Roman 12pt, double spacing, 1" margins,
    hanging indent for references.
    """
    from docx import Document                                    # deferred: not all routes need it
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

    doc = Document()

    # ── Page margins: 1" all sides ──────────────────────────────────
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    # ── Default style: Times New Roman 12pt ─────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    def para(text: str = "", bold: bool = False,
             align=WD_ALIGN_PARAGRAPH.LEFT,
             first_indent: float = 0.0,
             left_indent: float = 0.0) -> None:
        p = doc.add_paragraph()
        p.alignment = align
        fmt = p.paragraph_format
        fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        if first_indent:
            fmt.first_line_indent = Inches(first_indent)
        if left_indent:
            fmt.left_indent = Inches(left_indent)
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    # ── Title ────────────────────────────────────────────────────────
    para(review.get("title", "Literature Review"),
         bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ── Abstract ─────────────────────────────────────────────────────
    if review.get("abstract"):
        para("Abstract", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        para(review["abstract"])

    # ── Body sections ────────────────────────────────────────────────
    for section in review.get("sections", []):
        heading = section.get("heading", "")
        content = section.get("content", "")
        # APA Level 1 heading: centered, bold
        para(heading, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        # Body paragraph: 0.5" first-line indent
        para(content, first_indent=0.5)

    # ── References ───────────────────────────────────────────────────
    para("References", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ref in review.get("references", []):
        # APA hanging indent: 0.5" left, -0.5" first line
        para(ref, left_indent=0.5, first_indent=-0.5)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


router = APIRouter(tags=["graph"])


# ---------------------------------------------------------------------------
# REST — full graph snapshot
# ---------------------------------------------------------------------------

@router.get("/graph", response_model=GraphResponse)
async def get_graph(request: Request) -> GraphResponse:
    """Return the full knowledge graph for the initial frontend load.

    After this initial fetch the frontend relies exclusively on
    WebSocket delta events — it should never re-fetch /graph.
    """
    graph_manager = request.app.state.graph_manager
    snapshot = await graph_manager.get_full_graph()

    # Separate OpenQuestion nodes from regular nodes
    questions: list[dict] = []
    nodes: list[dict] = []
    for node in snapshot.get("nodes", []):
        if node.get("label") == "OpenQuestion":
            questions.append(node)
        else:
            nodes.append(node)

    return GraphResponse(
        nodes=nodes,
        edges=snapshot.get("edges", []),
        questions=questions,
    )


# ---------------------------------------------------------------------------
# REST — per-paper pipeline status
# ---------------------------------------------------------------------------

@router.get("/paper/{paper_id}/status", response_model=PaperStatusResponse)
async def get_paper_status(request: Request, paper_id: str) -> PaperStatusResponse:
    """Return the current pipeline status for a given paper."""
    graph_manager = request.app.state.graph_manager
    paper = await graph_manager.get_paper(paper_id)

    if paper is None:
        raise HTTPException(status_code=404, detail=f"Paper {paper_id} not found.")

    status = paper.get("status", "unknown")
    progress = PIPELINE_PROGRESS.get(status, 0.0)

    return PaperStatusResponse(
        paper_id=paper_id,
        status=status,
        progress_pct=progress,
    )


# ---------------------------------------------------------------------------
# REST — edge detail
# ---------------------------------------------------------------------------

@router.get("/edge/{edge_id}", response_model=EdgeDetailResponse)
async def get_edge_detail(
    edge_id: str,
    request: Request,
) -> EdgeDetailResponse:
    """Fetch full details for a relationship: both claims + their papers."""
    graph_manager = request.app.state.graph_manager
    result = await graph_manager.get_edge_with_claims(edge_id)
    if result is None:
        raise HTTPException(status_code=404, detail=f"Edge {edge_id!r} not found")
    return EdgeDetailResponse(**result)


# ---------------------------------------------------------------------------
# REST — consensus explanation
# ---------------------------------------------------------------------------

@router.get(
    "/claim/{claim_id}/consensus/explain",
    response_model=ConsensusExplainResponse,
)
async def explain_claim_consensus(
    claim_id: str,
    request: Request,
) -> ConsensusExplainResponse:
    """Generate an AI explanation of field consensus for a given claim."""
    graph_manager = request.app.state.graph_manager
    explainer = request.app.state.consensus_explainer

    context = await graph_manager.get_claim_consensus_context(claim_id)
    if context is None:
        raise HTTPException(status_code=404, detail=f"Claim {claim_id!r} not found")

    claim = context["claim"]
    semantic_edges = context["semantic_edges"]

    # Compute consensus_pct server-side (same formula as frontend)
    support_score = sum(
        (e.get("strength") or 0.5) * (0.5 if e["edge_type"] in ("EXTENDS", "REFINES") else 1.0)
        for e in semantic_edges if e["edge_type"] != "CONTRADICTS"
    )
    dispute_score = sum(
        (e.get("strength") or 0.5)
        for e in semantic_edges if e["edge_type"] == "CONTRADICTS"
    )
    total = support_score + dispute_score
    consensus_pct = round((support_score / total) * 100) if total > 0 else None

    explanation = await explainer.explain(
        claim_text=claim["text"],
        paper_title=claim["paper_title"],
        consensus_pct=consensus_pct or 0,
        semantic_edges=semantic_edges,
    )

    return ConsensusExplainResponse(
        claim_id=claim_id,
        explanation=explanation,
        consensus_pct=consensus_pct,
    )


# ---------------------------------------------------------------------------
# REST — claim verification
# ---------------------------------------------------------------------------

@router.get("/verify", response_model=VerificationResponse)
async def verify_statement(
    request: Request,
    statement: str = Query(..., min_length=5, max_length=500,
                           description="Statement to verify against the corpus"),
) -> VerificationResponse:
    """Verify a user-supplied statement against all claims in the corpus.

    Embeds the statement, pre-filters by cosine similarity, then classifies
    each candidate claim as supporting, contradicting, or neutral in a single
    batched LLM call.
    """
    graph_manager = request.app.state.graph_manager
    verifier = request.app.state.claim_verifier

    claims = await graph_manager.get_claims_for_verification()

    if not claims:
        return VerificationResponse(
            statement=statement,
            supports=[], contradicts=[], neutral=[],
            total_claims_checked=0,
        )

    result = await verifier.verify(statement=statement, claims=claims)

    return VerificationResponse(
        statement=statement,
        supports=[VerificationResultItem(**r) for r in result["supports"]],
        contradicts=[VerificationResultItem(**r) for r in result["contradicts"]],
        neutral=[VerificationResultItem(**r) for r in result["neutral"]],
        total_claims_checked=result["total_claims_checked"],
    )


# ---------------------------------------------------------------------------
# REST — research thread tracer
# ---------------------------------------------------------------------------

@router.get("/trace", response_model=TraceResponse)
async def trace_thread(
    request: Request,
    from_id: str = Query(..., min_length=1, description="Start node ID"),
    to_id: str = Query(..., min_length=1, description="End node ID"),
) -> TraceResponse:
    """Find the shortest path between two graph nodes and narrate it."""
    if from_id == to_id:
        raise HTTPException(400, "from_id and to_id must be different nodes")

    graph_manager = request.app.state.graph_manager
    tracer = request.app.state.thread_tracer

    path_data = await graph_manager.get_path_between(from_id, to_id)
    if not path_data:
        raise HTTPException(
            404,
            f"No path found between {from_id!r} and {to_id!r} within 8 hops",
        )

    narrative = await tracer.trace(
        path_nodes=path_data["nodes"],
        path_edges=path_data["edges"],
    )

    return TraceResponse(
        from_id=from_id,
        to_id=to_id,
        path_nodes=[TracePathNode(**n) for n in path_data["nodes"]],
        path_edges=[TracePathEdge(**e) for e in path_data["edges"]],
        path_length=path_data["path_length"],
        narrative=narrative,
    )


# ---------------------------------------------------------------------------
# REST — literature review generator
# ---------------------------------------------------------------------------

@router.post("/generate/review", response_model=LiteratureReviewResponse)
async def generate_review(request: Request) -> LiteratureReviewResponse:
    """Generate a full APA-format literature review from the loaded corpus."""
    graph_manager = request.app.state.graph_manager
    agent = request.app.state.literature_review

    review_data = await graph_manager.get_review_data()

    if not review_data.get("papers"):
        raise HTTPException(
            400,
            "No complete papers in corpus. Upload and process papers before generating a review.",
        )

    review_dict = await agent.generate(review_data)

    return LiteratureReviewResponse(
        review=LiteratureReviewData(**review_dict),
        paper_count=len(review_data["papers"]),
    )


@router.post("/generate/review/docx")
async def download_review_docx(payload: ReviewDownloadRequest) -> Response:
    """Convert a review JSON object to APA-formatted .docx for download."""
    try:
        docx_bytes = build_apa_docx(payload.review)
    except ImportError:
        raise HTTPException(
            500,
            "python-docx is not installed. Add python-docx==1.1.2 to requirements.txt and run pip install.",
        )

    title = str(payload.review.get("title", "literature_review"))
    filename = title.lower().replace(" ", "_")[:50].strip("_") + ".docx"

    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# WebSocket — live graph delta stream
# ---------------------------------------------------------------------------

@router.websocket("/ws/graph")
async def ws_graph(websocket: WebSocket) -> None:
    """Subscribe a client to real-time graph delta events.

    The server sends deltas; the client does not send meaningful messages
    but we keep the receive loop alive to detect disconnects.
    """
    emitter = websocket.app.state.emitter

    await websocket.accept()
    await emitter.subscribe(websocket)
    logger.info(
        "WebSocket client connected",
        extra={"subscriber_count": emitter.subscriber_count},
    )

    try:
        # Keep the connection alive — wait for the client to close it
        while True:
            # recv keeps the task alive; we ignore incoming data
            await websocket.receive_text()
    except WebSocketDisconnect:
        logger.info("WebSocket client disconnected normally")
    except Exception:
        logger.exception("WebSocket error")
    finally:
        await emitter.unsubscribe(websocket)
        logger.info(
            "WebSocket client unsubscribed",
            extra={"subscriber_count": emitter.subscriber_count},
        )
